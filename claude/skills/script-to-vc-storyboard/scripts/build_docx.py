"""Write Claude's shot plan (shots.json) as a Word storyboard table that Video Console imports one row per shot.

    python build_docx.py <shots.json> [--out <file.docx>]

Default output: <shots.json folder>/<source script name>-VC分鏡表.docx. Each row gets the fields the way the VC
breakdown writes them: name "場景（景別）", picture "景別：畫面", camera "景別，運鏡", speaker as the first notes
line "說話者：X". Every row has its own length, so VC keeps it as one shot.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

# (header cell, VC field). The header names are VC script-import aliases; check_import.py verifies them.
COLUMNS = (("鏡號", "index"), ("鏡頭名稱", "name"), ("畫面", "prompt"), ("動作", "motion"), ("運鏡", "camera"),
           ("字幕", "caption"), ("秒數", "duration"), ("備註", "notes"))
WIDTHS_CM = (1.0, 2.8, 7.0, 3.4, 2.6, 4.3, 1.2, 3.4)
FONT = "Microsoft JhengHei"
MIN_DURATION, MAX_DURATION = 0.5, 60.0
_EDGE = " \t\n，,；;：:、。．"


def _clean(value: object) -> str:
    """Collapse spaces and drop the edge punctuation VC trims anyway, so the table shows what VC imports."""
    lines = [re.sub(r"[ \t　]+", " ", ln).strip() for ln in str(value or "").split("\n")]
    return "\n".join(ln for ln in lines if ln).strip(_EDGE)


def _with_looks(picture: str, looks: dict[str, str]) -> str:
    """Each character's look after their first mention, unless the picture already describes it there."""
    if not looks:
        return picture
    seen: set[str] = set()

    def add(m: re.Match) -> str:
        name = m.group()
        if name in seen or picture.startswith(("（", "("), m.end()):
            return name
        seen.add(name)
        return f"{name}（{looks[name]}）"

    return re.compile("|".join(sorted(map(re.escape, looks), key=len, reverse=True))).sub(add, picture)


def rows(data: dict) -> list[dict]:
    """The table rows (VC field -> cell text) for shots.json; ValueError names every row that cannot be built."""
    looks = {str(k): _clean(v) for k, v in (data.get("characters") or {}).items() if _clean(v)}
    out, problems = [], []
    for i, shot in enumerate(data.get("shots") or [], 1):
        size, picture, scene = _clean(shot.get("size")), _clean(shot.get("picture")), _clean(shot.get("scene"))
        duration = shot.get("duration")
        if not size or not picture or not scene:
            problems.append(f"第 {i} 鏡缺少 scene、size 或 picture")
            continue
        if not isinstance(duration, (int, float)) or not MIN_DURATION <= duration <= MAX_DURATION:
            problems.append(f"第 {i} 鏡的 duration 要是 {MIN_DURATION:g} 到 {MAX_DURATION:g} 之間的數字")
            continue
        move = _clean(shot.get("move"))
        speaker = _clean(shot.get("speaker"))
        notes = ([f"說話者：{speaker}"] if speaker else []) + [_clean(shot.get("notes"))]  # VC's own order
        out.append({
            "act": _clean(shot.get("act")),
            "size": size,
            "index": str(i),
            "name": f"{scene}（{size}）",
            "prompt": f"{size}：{_with_looks(picture, looks)}",
            "motion": _clean(shot.get("motion")),
            "camera": f"{size}，{move}" if move else size,
            "caption": _clean(shot.get("caption")).strip("「」『』“”\"").rstrip("。．."),
            "duration": f"{round(float(duration), 1):g}",
            "notes": "\n".join(n for n in notes if n),
        })
    if problems:
        raise ValueError("\n".join(problems))
    return out


def default_out(json_path: Path, data: dict) -> Path:
    stem = Path(str(data.get("source") or json_path.stem)).stem
    return json_path.parent / f"{stem}-VC分鏡表.docx"


def _font(run, size: float, bold: bool = False) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def _fill(cell, text: str, size: float, bold: bool = False) -> None:
    """One Word paragraph per line, so VC reads the cell's line breaks back."""
    lines = text.split("\n") if text else [""]
    first = cell.paragraphs[0]
    for n, line in enumerate(lines):
        p = first if n == 0 else cell.add_paragraph()
        _font(p.add_run(line), size, bold)


def _table(doc, items: list[dict]):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm

    table = doc.add_table(rows=1, cols=len(COLUMNS))
    table.style = "Table Grid"
    table.autofit = False
    head = table.rows[0]
    flag = OxmlElement("w:tblHeader")  # repeat the header row on every page
    flag.set(qn("w:val"), "true")
    head._tr.get_or_add_trPr().append(flag)
    for cell, (label, _), width in zip(head.cells, COLUMNS, WIDTHS_CM):
        cell.width = Cm(width)
        _fill(cell, label, 9.5, bold=True)
    for item in items:
        cells = table.add_row().cells
        for cell, (_, field), width in zip(cells, COLUMNS, WIDTHS_CM):
            cell.width = Cm(width)
            _fill(cell, item[field], 9)
    return table


def build(data: dict, out: Path) -> list[dict]:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm

    table_rows = rows(data)
    if not table_rows:
        raise ValueError("shots.json 沒有任何鏡頭")
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.4))
    total = sum(float(r["duration"]) for r in table_rows)
    _font(doc.add_paragraph().add_run(_clean(data.get("title")) or "VC 分鏡表"), 16, bold=True)
    info = f"共 {len(table_rows)} 鏡，總長 {total:.1f} 秒"
    if data.get("source"):
        info += f"；依據：{Path(str(data['source'])).name}"
    _font(doc.add_paragraph().add_run(info), 10)
    for name, look in (data.get("characters") or {}).items():
        _font(doc.add_paragraph().add_run(f"角色外觀｜{name}：{_clean(look)}"), 9)
    groups: list[tuple[str, list[dict]]] = []
    for r in table_rows:
        if not groups or (r["act"] and r["act"] != groups[-1][0]):
            groups.append((r["act"], []))
        groups[-1][1].append(r)
    for act, items in groups:
        if act:
            _font(doc.add_paragraph().add_run(f"【{act}】"), 11, bold=True)
        _table(doc, items)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return table_rows


def main() -> int:
    ap = argparse.ArgumentParser(description="把 shots.json 寫成 VC 匯入用的 Word 分鏡表")
    ap.add_argument("shots", type=Path)
    ap.add_argument("--out", type=Path)
    args = ap.parse_args()
    data = json.loads(args.shots.read_text(encoding="utf-8"))
    out = args.out or default_out(args.shots, data)
    try:
        table_rows = build(data, out)
    except ValueError as exc:
        print(f"無法產生分鏡表：\n{exc}", file=sys.stderr)
        return 1
    print(f"已寫出 {out}（{len(table_rows)} 鏡）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
